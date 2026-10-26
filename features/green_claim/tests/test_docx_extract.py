"""Word(.docx) 텍스트 추출 검증. 외부 요청·OpenAI 호출 없음. 실행(저장소 루트):

    python -m unittest features.green_claim.tests.test_docx_extract -v

테스트 문서는 python-docx로 메모리에서 만드는 가상 문서입니다. API로 만들 수 없는 요소(하이퍼링크·텍스트 상자·
변경 추적·콘텐츠 컨트롤·각주 참조)는 XML을 직접 넣어 만듭니다. 실제 추출은 별도 프로세스(spawn)로도 실행합니다.
"""
import io
import multiprocessing
import struct
import unittest
import zipfile
import zlib
from unittest import mock

import docx
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.oxml.parser import OxmlElement
from werkzeug.datastructures import FileStorage

from app import create_app
from features.green_claim import docx_worker, service

KO = "리젠 T2T는 폐섬유를 재활용한 원료로 만듭니다."
EN = "Made with recycled textile waste."
PT = "Fabricado com resíduos têxteis reciclados. Ação, coração, não."
LIMITS = service.docx_limits()


def to_bytes(document):
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def tiny_png():
    def chunk(kind, data):
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)
    raw = zlib.compress(b"\x00\xff\xff\xff")
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", raw) + chunk(b"IEND", b""))


def add_hyperlink(paragraph, text, url):
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def rewrite_zip(data, extra=(), replace=None, compression=zipfile.ZIP_DEFLATED):
    """ZIP 항목을 다시 써서 항목 추가·내용 교체."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            content = src.read(info.filename)
            if replace and info.filename in replace:
                content = replace[info.filename](content)
            dst.writestr(info.filename, content)
        for name, content in extra:
            dst.writestr(name, content, compress_type=compression)
    return out.getvalue()


def base_document():
    d = docx.Document()
    d.add_paragraph("PARA-1 " + KO)
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "EN", EN
    t.cell(1, 0).text, t.cell(1, 1).text = "PT", PT
    d.add_paragraph("PARA-2 depois da tabela")
    return d


class WorkerTest(unittest.TestCase):
    """docx_worker.extract를 현재 프로세스에서 직접 확인(실제 python-docx)."""

    def test_order_languages_and_ranges(self):
        r = docx_worker.extract(to_bytes(base_document()), LIMITS)
        self.assertEqual(r["status"], "ok")
        text = r["text"]
        self.assertTrue(text.index("PARA-1") < text.index(EN) < text.index(PT) < text.index("PARA-2"))
        self.assertIn(KO, text)
        self.assertEqual([b["type"] for b in r["blocks"]], ["paragraph", "table", "paragraph"])
        table = r["blocks"][1]
        self.assertEqual(text[table["char_start"]:table["char_end"]], f"EN | {EN}\nPT | {PT}")
        cell = table["rows"][1]["cells"][1]
        self.assertEqual(text[cell["char_start"]:cell["char_end"]], PT)
        self.assertEqual(r["detected_unsupported"], [])
        # 블록 사이 빈 줄 포함 범위가 이어짐
        self.assertEqual(table["char_start"], r["blocks"][0]["char_end"] + 2)

    def test_merged_cells_not_repeated(self):
        d = docx.Document()
        t = d.add_table(rows=3, cols=3)
        t.cell(0, 0).merge(t.cell(0, 1)).text = "WIDE"
        t.cell(0, 2).text = "C"
        t.cell(1, 0).merge(t.cell(2, 0)).text = "TALL"
        t.cell(1, 1).text, t.cell(1, 2).text = "b1", "c1"
        t.cell(2, 1).text, t.cell(2, 2).text = "b2", "c2"
        r = docx_worker.extract(to_bytes(d), LIMITS)
        self.assertEqual(r["text"], "WIDE | C\nTALL | b1 | c1\n | b2 | c2")
        self.assertEqual(r["text"].count("WIDE"), 1)
        self.assertEqual(r["text"].count("TALL"), 1)
        last_row = r["blocks"][0]["rows"][2]["cells"]
        self.assertTrue(last_row[0]["merged_continue"])

    def test_nested_table(self):
        d = docx.Document()
        outer = d.add_table(rows=1, cols=2)
        outer.cell(0, 0).text = "OUTER"
        inner = outer.cell(0, 1).add_table(rows=1, cols=2)
        inner.cell(0, 0).text, inner.cell(0, 1).text = "IN-A", "IN-B"
        r = docx_worker.extract(to_bytes(d), LIMITS)
        self.assertIn("OUTER", r["text"])
        self.assertIn("IN-A | IN-B", r["text"])
        self.assertEqual(r["nested_table_count"], 1)

    def test_hyperlink_visible_text(self):
        d = docx.Document()
        p = d.add_paragraph("Veja ")
        add_hyperlink(p, "nosso site", "https://example.com")
        p.add_run(" agora.")
        r = docx_worker.extract(to_bytes(d), LIMITS)
        self.assertEqual(r["text"], "Veja nosso site agora.")

    def test_detects_unsupported_elements_and_omits_their_text(self):
        d = base_document()
        d.sections[0].header.paragraphs[0].text = "HEADER TEXT"
        d.add_picture(io.BytesIO(tiny_png()))
        p = d.add_paragraph("BEFORE ")
        p._p.append(parse_xml(
            f'<w:r {nsdecls("w")} xmlns:v="urn:schemas-microsoft-com:vml"><w:pict><v:shape><v:textbox>'
            f'<w:txbxContent><w:p><w:r><w:t>BOX TEXT</w:t>'
            f'</w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict></w:r>'))
        p._p.append(parse_xml(f'<w:ins {nsdecls("w")} w:id="1" w:author="x"><w:r><w:t>INSERTED</w:t></w:r></w:ins>'))
        p._p.append(parse_xml(f'<w:r {nsdecls("w")}><w:footnoteReference w:id="1"/></w:r>'))
        d.element.body.insert(0, parse_xml(
            f'<w:sdt {nsdecls("w")}><w:sdtContent><w:p><w:r><w:t>SDT TEXT</w:t></w:r></w:p></w:sdtContent></w:sdt>'))
        d.add_comment(d.paragraphs[0].runs, text="memo", author="tester")
        r = docx_worker.extract(to_bytes(d), LIMITS)
        codes = {f["code"] for f in r["detected_unsupported"]}
        self.assertTrue({"headers_footers", "images", "text_boxes", "tracked_changes", "notes",
                         "content_controls", "comments"} <= codes, codes)
        for missing in ("HEADER TEXT", "BOX TEXT", "INSERTED", "SDT TEXT", "memo"):
            self.assertNotIn(missing, r["text"])
        self.assertIn("BEFORE", r["text"])

    def test_image_only_and_empty(self):
        d = docx.Document()
        d.add_picture(io.BytesIO(tiny_png()))
        self.assertEqual(docx_worker.extract(to_bytes(d), LIMITS)["code"], "DOCX_NO_TEXT")
        self.assertEqual(docx_worker.extract(to_bytes(docx.Document()), LIMITS)["code"], "DOCX_NO_TEXT")

    def test_structure_errors(self):
        good = to_bytes(base_document())
        self.assertEqual(docx_worker.extract(good[: len(good) // 2], LIMITS)["code"], "DOCX_DAMAGED")
        other_zip = io.BytesIO()
        with zipfile.ZipFile(other_zip, "w") as z:
            z.writestr("readme.txt", "not word")
        self.assertEqual(docx_worker.extract(other_zip.getvalue(), LIMITS)["code"], "DOCX_INVALID")
        xlsx_like = rewrite_zip(good, replace={"[Content_Types].xml": lambda c: c.replace(
            b"wordprocessingml.document.main+xml", b"spreadsheetml.sheet.main+xml")})
        self.assertEqual(docx_worker.extract(xlsx_like, LIMITS)["code"], "DOCX_INVALID")
        docm = rewrite_zip(good, replace={"[Content_Types].xml": lambda c: c.replace(
            b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
            b"application/vnd.ms-word.document.macroEnabled.main+xml")})
        self.assertEqual(docx_worker.extract(docm, LIMITS)["code"], "DOCX_UNSUPPORTED_FORMAT")
        broken_xml = rewrite_zip(good, replace={"word/document.xml": lambda c: c[: len(c) // 2]})
        self.assertEqual(docx_worker.extract(broken_xml, LIMITS)["code"], "DOCX_DAMAGED")

    def test_external_entity_not_resolved(self):
        good = to_bytes(base_document())
        evil = rewrite_zip(good, replace={"[Content_Types].xml": lambda c: c.replace(
            b"<Types", b'<!DOCTYPE Types [<!ENTITY x SYSTEM "file:///C:/Windows/win.ini">]><Types', 1)})
        r = docx_worker.extract(evil, LIMITS)
        # 엔티티를 확장하지 않으므로 파일 내용이 결과에 들어가지 않음(성공 또는 손상 처리)
        self.assertNotIn("[fonts]", str(r))

    def test_zip_entry_limit_boundary(self):
        good = to_bytes(base_document())
        count = len(zipfile.ZipFile(io.BytesIO(good)).infolist())
        at_limit = rewrite_zip(good, extra=[(f"customXml/pad{i}.txt", b"x") for i in range(1000 - count)])
        self.assertEqual(len(zipfile.ZipFile(io.BytesIO(at_limit)).infolist()), 1000)
        self.assertEqual(docx_worker.extract(at_limit, LIMITS)["status"], "ok")
        over = rewrite_zip(good, extra=[(f"customXml/pad{i}.txt", b"x") for i in range(1001 - count)])
        self.assertEqual(docx_worker.extract(over, LIMITS)["code"], "DOCX_ZIP_TOO_MANY_ENTRIES")

    def test_uncompressed_size_boundary(self):
        good = to_bytes(base_document())
        used = sum(i.file_size for i in zipfile.ZipFile(io.BytesIO(good)).infolist())
        at_limit = rewrite_zip(good, extra=[("customXml/zeros.bin", b"\x00" * (50_000_000 - used))])
        self.assertLess(len(at_limit), 1_000_000)  # 압축 파일은 작지만 해제 크기는 상한
        self.assertEqual(docx_worker.extract(at_limit, LIMITS)["status"], "ok")
        over = rewrite_zip(good, extra=[("customXml/zeros.bin", b"\x00" * (50_000_001 - used))])
        self.assertEqual(docx_worker.extract(over, LIMITS)["code"], "DOCX_ZIP_TOO_LARGE")

    def test_actual_decompressed_bytes_counted(self):
        """헤더 선언 크기가 작아도 실제로 읽은 양으로 상한을 지키는지(모의 ZIP 객체)."""
        class Info:
            filename, file_size, flag_bits = "a", 0, 0

        class FakeZip:
            def infolist(self):
                return [Info()]

            def open(self, info):
                return io.BytesIO(b"\x00" * 2000)

        with self.assertRaises(docx_worker._Fail) as ctx:
            docx_worker._check_zip(FakeZip(), 1000, 1000)
        self.assertEqual(ctx.exception.code, "DOCX_ZIP_TOO_LARGE")

    def test_lying_header_rejected(self):
        """중앙 디렉터리·로컬 헤더의 해제 크기를 실제보다 작게 바꾼 ZIP은 선언 크기 이상 풀지 않고 손상 처리."""
        good = to_bytes(base_document())
        data = bytearray(rewrite_zip(good, extra=[("customXml/zeros.bin", b"\x00" * 3_000_000)]))
        zf = zipfile.ZipFile(io.BytesIO(bytes(data)))
        info = zf.getinfo("customXml/zeros.bin")
        struct.pack_into("<I", data, info.header_offset + 22, 10)  # 로컬 헤더 uncompressed size
        patched = 0
        pos = bytes(data).find(b"PK\x01\x02")
        while pos >= 0:
            name_len = struct.unpack_from("<H", data, pos + 28)[0]
            if bytes(data[pos + 46: pos + 46 + name_len]) == b"customXml/zeros.bin":
                struct.pack_into("<I", data, pos + 24, 10)  # 중앙 디렉터리 uncompressed size
                patched += 1
            pos = bytes(data).find(b"PK\x01\x02", pos + 4)
        self.assertEqual(patched, 1)
        self.assertEqual(zipfile.ZipFile(io.BytesIO(bytes(data))).getinfo("customXml/zeros.bin").file_size, 10)
        self.assertEqual(docx_worker.extract(bytes(data), LIMITS)["code"], "DOCX_DAMAGED")

    def test_char_limit_boundary_includes_separator(self):
        d = docx.Document()
        d.add_paragraph("a" * 9999)
        d.add_paragraph("b" * 9999)
        r = docx_worker.extract(to_bytes(d), LIMITS)
        self.assertEqual(len(r["text"]), 20000)
        self.assertEqual(service.build_docx_result(r)["char_count"], 20000)
        d.add_paragraph("c")
        self.assertEqual(docx_worker.extract(to_bytes(d), LIMITS)["code"], "TEXT_TOO_LONG")


class ServiceTest(unittest.TestCase):
    def test_signature_checks(self):
        ole = service.OLE_SIGNATURE + b"\x00" * 600
        with self.assertRaises(service.ExtractError) as ctx:
            service.validate_docx_upload("a.docx", ole + service.OLE_ENCRYPTED_PACKAGE)
        self.assertEqual(ctx.exception.code, "DOCX_ENCRYPTED")
        with self.assertRaises(service.ExtractError) as ctx:
            service.validate_docx_upload("a.docx", ole)
        self.assertEqual(ctx.exception.code, "DOCX_UNSUPPORTED_FORMAT")
        with self.assertRaises(service.ExtractError) as ctx:
            service.validate_docx_upload("a.docx", "그냥 텍스트".encode())
        self.assertEqual(ctx.exception.code, "DOCX_INVALID")
        for name in ("a.doc", "a.docm", "a.dotx", "a.rtf", "a.png"):
            with self.assertRaises(service.ExtractError) as ctx:
                service.extract_file(name, b"PK\x03\x04data")
            self.assertEqual(ctx.exception.code, "FILE_TYPE_UNSUPPORTED", name)

    def test_real_process_and_timeout(self):
        r = service.extract_file("promo.DOCX", to_bytes(base_document()))
        self.assertEqual(r["format"], "docx")
        self.assertNotIn("pages", r)
        self.assertEqual(multiprocessing.active_children(), [])
        with mock.patch.object(service.logger, "warning") as warn:
            with self.assertRaises(service.ExtractError) as ctx:
                service.extract_docx("a.docx", to_bytes(base_document()), timeout=0.05)
        self.assertEqual(ctx.exception.code, "EXTRACT_TIMEOUT")
        self.assertNotIn(KO, str(warn.call_args))
        self.assertEqual(multiprocessing.active_children(), [])

    def test_worker_exception_maps_to_failed(self):
        with mock.patch.object(docx_worker, "extract", side_effect=RuntimeError("boom")):
            in_recv, in_send = multiprocessing.Pipe(duplex=False)
            out_recv, out_send = multiprocessing.Pipe(duplex=False)
            in_send.send_bytes(b"x")
            docx_worker.run(in_recv, out_send, LIMITS)
            self.assertEqual(out_recv.recv(), {"status": "error", "code": "EXTRACT_FAILED"})


class RouteTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.client = create_app().test_client()

    def post(self, filename, content):
        return self.client.post("/api/green-claim/extract", content_type="multipart/form-data",
                                data={"promo_file": FileStorage(io.BytesIO(content), filename=filename)})

    def test_success_shape(self):
        d = base_document()
        d.add_picture(io.BytesIO(tiny_png()))
        resp = self.post("promo.docx", to_bytes(d))
        body = resp.get_json()
        self.assertEqual(resp.status_code, 200, body)
        data = body["data"]
        self.assertEqual(set(data), {"format", "text", "char_count", "block_count", "paragraph_count", "table_count",
                                     "blocks", "empty_block_count", "nested_table_count", "detected_unsupported",
                                     "notices", "limits"})
        self.assertEqual(data["format"], "docx")
        self.assertEqual(data["char_count"], len(data["text"]))
        self.assertIn("images", {f["code"] for f in data["detected_unsupported"]})
        self.assertEqual(data["limits"]["max_zip_entries"], 1000)

    def test_errors_and_size_boundary(self):
        for name, content, code, status in [
            ("a.docx", b"", "FILE_EMPTY", 400),
            ("a.doc", b"x", "FILE_TYPE_UNSUPPORTED", 415),
            ("a.docx", b"plain text", "DOCX_INVALID", 422),
        ]:
            resp = self.post(name, content)
            self.assertEqual((resp.status_code, resp.get_json()["error"]["code"]), (status, code))
        good = to_bytes(base_document())
        pad = 10_000_000 - len(good)
        for _ in range(5):
            data = rewrite_zip(good, extra=[("customXml/pad.bin", bytes(range(256)) * (pad // 256) + b"x" * (pad % 256))],
                               compression=zipfile.ZIP_STORED)
            if len(data) == 10_000_000:
                break
            pad += 10_000_000 - len(data)
        self.assertEqual(len(data), 10_000_000)
        self.assertEqual(self.post("big.docx", data).status_code, 200)
        resp = self.post("big.docx", data + b"\x00")
        self.assertEqual((resp.status_code, resp.get_json()["error"]["code"]), (413, "FILE_TOO_LARGE"))


if __name__ == "__main__":
    unittest.main()
